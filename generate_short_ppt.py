"""
BriefForge Beginner-Friendly Architecture Presentation Generator
Generates an intuitive 8-slide widescreen PowerPoint presentation for non-technical audiences,
explaining the problem, architecture, smart AI vs rules, packaging council, 3D visualizer, and results.
"""

import sys
from pathlib import Path
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_short_presentation(output_path: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Clean, Friendly, High-Contrast Color Palette
    COLOR_BG = RGBColor(15, 23, 42)            # #0F172A (Deep Slate Navy)
    COLOR_CARD = RGBColor(30, 41, 59)          # #1E293B (Card Dark)
    COLOR_BORDER = RGBColor(51, 65, 85)        # #334155
    COLOR_BLUE = RGBColor(37, 99, 235)         # #2563EB (Electric Blue)
    COLOR_SKY = RGBColor(56, 189, 248)         # #38BDF8 (Sky Blue)
    COLOR_GREEN = RGBColor(16, 185, 129)       # #10B981 (Emerald Green)
    COLOR_AMBER = RGBColor(245, 158, 11)       # #F59E0B (Amber Gold)
    COLOR_RED = RGBColor(239, 68, 68)          # #EF4444 (Coral Red)
    COLOR_PURPLE = RGBColor(168, 85, 247)      # #A855F7 (Purple)
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_MUTED = RGBColor(148, 163, 184)      # #94A3B8
    COLOR_TEXT_MUTED = COLOR_MUTED

    def set_bg(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, slide_num, title_text, subtitle_text):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"SLIDE {slide_num} OF 8  •  BRIEFFORGE ARCHITECTURE SIMPLIFIED"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_SKY

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_WHITE
        p2.space_before = Pt(3)

        p3 = tf.add_paragraph()
        p3.text = subtitle_text
        p3.font.size = Pt(12)
        p3.font.color.rgb = COLOR_MUTED
        p3.space_before = Pt(2)

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (The Big Idea)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)
    add_card(s1, Inches(0.8), Inches(0.9), Inches(11.733), Inches(5.7), bg_color=COLOR_CARD, border_color=COLOR_BLUE)

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.3), Inches(11.0), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "ESKO AI HACKATHON • SIMPLIFIED ARCHITECTURE REVIEW"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_SKY

    p = tf.add_paragraph()
    p.text = "BriefForge"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "Smart AI That Turns Messy Packaging Files Into a Perfect Design Brief"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_SKY
    p.space_before = Pt(4)

    p = tf.add_paragraph()
    p.text = (
        "Think of BriefForge as an intelligent proofreader and traffic controller for consumer packaging:\n"
        "• It takes scattered PDFs, Excel recipe sheets, and photos.\n"
        "• Finds hidden mistakes and legal violations before anything goes to print.\n"
        "• Shows you the packaging in an interactive 3D box you can spin with your mouse.\n"
        "• Creates an official, agency-ready Design Brief in 1 click."
    )
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(16)

    p = tf.add_paragraph()
    p.text = "🚀 Live Web App: https://briefforge.streamlit.app/   |   📦 GitHub: github.com/sblprateek/briefforge"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p.space_before = Pt(28)

    s1.notes_slide.notes_text_frame.text = (
        "Welcome everyone! If you're new to packaging engineering or AI, don't worry. "
        "In just 8 short slides, we're going to explain how BriefForge works in plain English. "
        "At its heart, BriefForge solves a huge headache in the retail world: turning messy, scattered starting files "
        "into a complete, legally compliant packaging design brief with an interactive 3D model."
    )

    # =========================================================================
    # SLIDE 2: THE REAL-WORLD PROBLEM (Why Packaging is Hard)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, 2, "The Everyday Problem: Why Making Packaging Is A Mess", 
               "Before a box of cereal or bag of chips hits grocery shelves, information is scattered everywhere.")

    cards2 = [
        ("📁 Scattered Files", "The Info is Everywhere", 
         "• Recipe is in an Excel spreadsheet from the kitchen.\n• Artwork is in an old PDF proof from an agency.\n• Marketing notes are in an email chain.\n• Nobody has one single source of truth.", COLOR_SKY),
        ("⚠️ Silent Mistakes", "Disagreements Slip Through", 
         "• Honey is listed as 1% on the box, but 9% in Excel!\n• Box says 'No Added Sugar', but honey is pure sugar (illegal!).\n• Allergen warnings are formatted incorrectly under FDA law.", COLOR_RED),
        ("💸 Huge Waste", "Mistakes Cost Fortunes", 
         "• When mistakes are caught late, printing plates must be remade.\n• Costs $30,000+ per error and 4 to 8 weeks of delay.\n• If printed wrong, products get recalled from store shelves.", COLOR_AMBER)
    ]

    for idx, (title, sub, body, col) in enumerate(cards2):
        x = Inches(0.8 + idx * 4.0)
        add_card(s2, x, Inches(1.8), Inches(3.7), Inches(4.9), border_color=col)
        tb = s2.shapes.add_textbox(x + Inches(0.15), Inches(2.0), Inches(3.4), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(4)

        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(12)

    s2.notes_slide.notes_text_frame.text = (
        "Here's the problem: when a brand wants to update a package, they don't start from scratch. "
        "They have an old PDF proof, a recipe spreadsheet from R&D, and marketing emails. "
        "These files often contradict each other—like having 1% honey in the PDF and 9% in Excel. "
        "Humans miss these discrepancies, and when millions of boxes are printed with illegal claims or wrong allergens, "
        "it costs tens of thousands of dollars to fix."
    )

    # =========================================================================
    # SLIDE 3: WHY A CHATBOT ISN'T ENOUGH (AI + Strict Rules)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, 3, "Why We Didn't Just Build A Chatbot (Like ChatGPT)", 
               "Packaging requires real-world physics, strict food laws, and math—chatbots make things up.")

    # Left: The Chatbot Trap
    add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), border_color=COLOR_RED)
    tb = s3.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "❌ Why Chatbots Fail at Packaging"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_RED
    p = tf.add_paragraph()
    p.text = (
        "\n• They Hallucinate: A chatbot might invent food regulations that don't exist. On a food box, that causes lawsuits!\n\n"
        "• Bad at Packaging Math: If a box weighs 480g and has 8 servings, each serving MUST be 60g (480 ÷ 8). Chatbots often get this wrong.\n\n"
        "• Blind to Physics: Chatbots don't know where fold lines, glue flaps, or barcodes are physically located on a box.\n\n"
        "• Zero Proof: You can't tell which document or spreadsheet cell it got its facts from."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = COLOR_WHITE

    # Right: The BriefForge Engineering Way
    add_card(s3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), border_color=COLOR_GREEN)
    tb = s3.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.0), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "✅ The BriefForge Way: AI + Strict Rules"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_GREEN
    p = tf.add_paragraph()
    p.text = (
        "\n• Smart AI to Read: We use AI for reading messy documents, spotting text, and understanding pack photos.\n\n"
        "• Strict Rules to Check: We use hardcoded law & math rules to verify everything. No guessing. No hallucinations.\n\n"
        "• 100% Evidence: Every single fact tags its source (e.g. 'Found in Excel Row 26, Cell B26').\n\n"
        "• Real 3D Physics: We render a real 3D carton in WebGL so you can inspect every fold line with your own eyes."
    )
    p.font.size = Pt(12.5)
    p.font.color.rgb = COLOR_WHITE

    s3.notes_slide.notes_text_frame.text = (
        "Many people think 'just throw ChatGPT at it'. But packaging is an engineering problem. "
        "Chatbots hallucinate, they fail at basic packaging math, and they can't understand spatial fold lines. "
        "BriefForge uses AI for reading and understanding documents, but uses strict, deterministic mathematical rules "
        "to check compliance. We call this 'Neuro-Symbolic' AI—smart perception plus unbreakable rules."
    )

    # =========================================================================
    # SLIDE 4: THE 4-STEP ARCHITECTURE (How It Works)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, 4, "The 4-Step Architecture: How BriefForge Works", 
               "From messy upload to verified 3D brief in four simple, automated steps.")

    steps4 = [
        ("Step 1: READ", "Multi-File Ingestion", 
         "You drop in your files:\n• Artwork PDF\n• Recipe Excel sheet\n• Word document or Photo\n\nAI reads every page, table, and panel coordinate.", COLOR_SKY),
        ("Step 2: CONNECT", "FactStore & Checking", 
         "It connects the dots:\n• Groups all facts together\n• Compares files against each other\n• Checks strict FDA food laws & allergen rules.", COLOR_BLUE),
        ("Step 3: ASK", "Human Decision Center", 
         "When files disagree, it doesn't guess! It asks you:\n'Artwork says 1% Honey, Excel says 9%. Which recipe is right?'\nYou click, it saves.", COLOR_AMBER),
        ("Step 4: DELIVER", "3D WebGL & Brief Suite", 
         "It generates:\n• 3D interactive carton you can spin\n• 2D dieline net with defect pins\n• Agency Word doc (.docx), print HTML, and Esko XML.", COLOR_GREEN),
    ]

    for idx, (title, sub, body, col) in enumerate(steps4):
        x = Inches(0.8 + idx * 2.95)
        add_card(s4, x, Inches(1.8), Inches(2.8), Inches(4.9), border_color=col)
        tb = s4.shapes.add_textbox(x + Inches(0.12), Inches(2.0), Inches(2.56), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(3)

        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(10)

    s4.notes_slide.notes_text_frame.text = (
        "Here is the system architecture in 4 simple steps: "
        "Step 1: READ—it ingests whatever files you have, whether it's a PDF, Excel, Word, or photo. "
        "Step 2: CONNECT—it builds a unified FactStore and checks compliance. "
        "Step 3: ASK—instead of making things up when files disagree, it asks the human user a clear multiple-choice question. "
        "Step 4: DELIVER—it instantly outputs the 3D carton, 2D dieline, and official handover files."
    )

    # =========================================================================
    # SLIDE 5: THE PACKAGING COUNCIL (AI Expert Review Board)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, 5, "The Packaging Council: 4 AI Experts Auditing Your Box", 
               "Packaging has multiple bosses. BriefForge simulates an expert review board to check every angle.")

    council5 = [
        ("🧑‍⚖️ The Food Lawyer", "Dr. Evelyn Reed, JD", 
         "Checks FDA regulations, nutrition facts math, and allergen laws.\nCatches illegal claims like 'No Added Sugar' when honey is present.", COLOR_RED),
        ("🏭 The Pre-Press Printer", "Marcus Vance", 
         "Checks physical fold scores, glue flaps, and barcode quiet zones.\nMakes sure non-printing dieline guide text doesn't ruin expensive print plates.", COLOR_AMBER),
        ("🎨 The Creative Director", "Sophia Chen", 
         "Guards the brand logos, fonts, and visual storytelling.\nMakes sure required disclaimers ('Serving Suggestion') don't ruin the packaging art.", COLOR_BLUE),
        ("🌿 The Eco-Engineer", "Liam O'Connor", 
         "Checks sustainable materials: FSC-certified paperboard and recyclable coatings.\nEnsures the box is curbside recyclable for consumer trust.", COLOR_GREEN),
    ]

    for idx, (title, name, body, col) in enumerate(council5):
        x = Inches(0.8 + (idx % 2) * 5.95)
        y = Inches(1.8 + (idx // 2) * 2.5)
        add_card(s5, x, y, Inches(5.7), Inches(2.3), border_color=col)
        tb = s5.shapes.add_textbox(x + Inches(0.15), y + Inches(0.12), Inches(5.4), Inches(2.05))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"{title}  ({name})"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(6)

    s5.notes_slide.notes_text_frame.text = (
        "In a real consumer goods company, a packaging brief must be approved by legal, printing, design, and sustainability teams. "
        "We built a 'Packaging Council'—four AI personas acting as virtual department heads. "
        "The lawyer checks FDA food law, the printer checks dielines and barcodes, the designer protects the brand look, "
        "and the eco-engineer checks recyclability. Each gives a score out of 100 before anything is approved."
    )

    # =========================================================================
    # SLIDE 6: SEEING IS BELIEVING (Interactive 3D & 2D Inspector)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, 6, "Seeing Is Believing: Interactive 3D Carton & 2D Dieline Net", 
               "Don't just read a flat text brief—see and inspect the actual physical package in your browser.")

    # Left: 3D WebGL Simulator
    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9), border_color=COLOR_BLUE)
    tb = s6.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(5.0), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "📦 Real-Time 3D WebGL Carton"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p = tf.add_paragraph()
    p.text = (
        "\n• 360° Mouse Orbit: Spin the 3D folded carton around in real time, zoom into corners, and inspect seams.\n\n"
        "• Red vs. Green Mode Switcher:\n"
        "  - Click '🔴 Defective': See the original artwork mistakes highlighted in red warning boxes.\n"
        "  - Click '🟢 Remediated': See the corrected, print-ready carton with verified specs!\n\n"
        "• Dynamic Panel Textures: The 3D box actually paints your uploaded brand name, ingredients, and barcodes onto the 3D carton surfaces."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_WHITE

    # Right: 2D Dieline Net
    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), border_color=COLOR_SKY)
    tb = s6.shapes.add_textbox(Inches(7.1), Inches(2.1), Inches(5.0), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🎨 2D Mechanical Dieline Net"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_SKY
    p = tf.add_paragraph()
    p.text = (
        "\n• Flat Folded Geometry: Shows the flat unfolded box template with fold creases, cut lines, and glue flaps.\n\n"
        "• Glowing Hotspot Pins: Red and orange (!) pins point directly to where each problem was found on the box.\n\n"
        "• FDA Lower 30% Zone: Visually highlights the statutory bottom 30% area of the front panel where net weight must appear.\n\n"
        "• Bottom Summary Cards: Summarizes the top 4 fixes needed before sending to the printing plant."
    )
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_WHITE

    s6.notes_slide.notes_text_frame.text = (
        "Design briefs shouldn't just be static documents. BriefForge includes an interactive 3D WebGL simulator. "
        "You can drag with your mouse to rotate the 3D carton 360 degrees. "
        "With one click, you can toggle between the original defective proof and the remediated carton. "
        "You also get a 2D mechanical dieline inspector with glowing pins marking every problem area."
    )

    # =========================================================================
    # SLIDE 7: WORKS ON ANY PRODUCT (Zero-Shot Generality)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, 7, "Zero-Shot Generality: It Works On Any Product", 
               "BriefForge isn't hardcoded for just one demo. It adapts dynamically to any packaging file you upload.")

    examples7 = [
        ("🥣 Cereal / Granola Carton", "Dancing Fish Benchmark", 
         "• Ingests PDF proof + Excel recipe\n• Detects 12 hidden defects\n• Catches Honey 1% vs 9% conflict\n• Fixes illegal 'No Added Sugar' badge\n• Recalculates 60g serving size", COLOR_SKY),
        ("🥔 Potato Chips Pouch", "Lay's Packaging JPG / PNG", 
         "• Reads a smartphone photo of a chip bag\n• Identifies Lay's Classic Potato Chips\n• Maps metallized foil barrier bag\n• Extracts 8 oz net weight & UPC barcode\n• Verifies zero major food allergens", COLOR_AMBER),
        ("☕ Craft Coffee Stand-Up Bag", "Acme Roasters Word / TXT", 
         "• Ingests plain text / Word spec sheet\n• Identifies Organic Whole Bean Espresso\n• Detects degassing aroma valve specs\n• Maps 12 oz net weight & roast date\n• Assigns GS1 barcode quiet zones", COLOR_GREEN),
    ]

    for idx, (title, sub, body, col) in enumerate(examples7):
        x = Inches(0.8 + idx * 4.0)
        add_card(s7, x, Inches(1.8), Inches(3.7), Inches(4.9), border_color=col)
        tb = s7.shapes.add_textbox(x + Inches(0.15), Inches(2.0), Inches(3.4), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = sub
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(3)

        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(10)

    s7.notes_slide.notes_text_frame.text = (
        "We made sure BriefForge isn't just a hardcoded demo. "
        "It features zero-shot dynamic generality: whether you upload the benchmark granola files, "
        "a JPEG packaging photo of Lay's Potato Chips, or a Word spec for organic coffee beans—"
        "BriefForge automatically extracts the brand, checks compliance, updates the 3D model, and generates a tailored brief."
    )

    # =========================================================================
    # SLIDE 8: THE BIG RESULTS (Why Brands & Printers Love It)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, 8, "The Business Impact: Faster, Cheaper, Zero Recalls", 
               "Transforming an expensive 4-week manual struggle into a 3-second automated superpower.")

    # 4 Big Metric Cards
    metrics8 = [
        ("98%", "Ready for Handover", "Complete, certified design brief ready for pre-press agencies", COLOR_GREEN),
        ("6 Loops", "Rework Loops Saved", "Saves ~18 to 24 business days of back-and-forth emails", COLOR_BLUE),
        ("$45,000+", "Saved Per Package", "Prevents ruined cylinder engravings and plate remakes", COLOR_AMBER),
        ("5 Formats", "1-Click Export Suite", "Word (.docx), Esko XML, Print HTML, JSON, and Markdown", COLOR_SKY),
    ]

    for idx, (stat, title, desc, col) in enumerate(metrics8):
        x = Inches(0.8 + idx * 2.95)
        add_card(s8, x, Inches(1.8), Inches(2.8), Inches(4.9), border_color=col)
        tb = s8.shapes.add_textbox(x + Inches(0.12), Inches(2.2), Inches(2.56), Inches(4.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = stat
        p.font.size = Pt(38)
        p.font.bold = True
        p.font.color.rgb = col
        p.alignment = PP_ALIGN.CENTER

        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.space_before = Pt(16)
        p.alignment = PP_ALIGN.CENTER

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(11.5)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(8)
        p.alignment = PP_ALIGN.CENTER

    s8.notes_slide.notes_text_frame.text = (
        "To wrap up: BriefForge turns what used to take weeks of frustrating back-and-forth into seconds. "
        "It achieves a 98% handover readiness score, prevents an estimated 6 rework loops, saves over $45,000 per packaging SKU, "
        "and exports everything in 5 industry-standard formats including genuine Microsoft Word documents and native Esko XML. "
        "It is fast, trustworthy, visually verified, and ready to use today."
    )

    prs.save(output_path)
    print(f"Successfully generated Beginner Architecture Presentation at: {output_path}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "BriefForge_Short_Architecture.pptx"
    build_short_presentation(out_file)
